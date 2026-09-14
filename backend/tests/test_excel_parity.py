"""Confronto con l'oracolo: l'app deve dire gli stessi numeri del workbook.

L'app non importa piu' da Money.xlsx e non ci scrive: questi test sono
l'ultimo legame, e sono un legame voluto. Leggono il workbook in sola lettura
e ricalcolano gli stessi aggregati che l'API espone, usando le formule
verificate a mano nel file (Budget Dashboard, Goals, colonna N di
Transactions). Se un numero diverge, e' l'app ad avere torto.

Richiedono il workbook al path indicato da MONEY_XLSX_PATH e un database
popolato con gli stessi dati: senza il file vengono saltati invece di fallire.
La copertura che non dipende da nulla di esterno sta in ``test_aggregates`` e
``test_interchange_roundtrip``.
"""
from __future__ import annotations

import os
from collections import defaultdict
from datetime import date
from pathlib import Path
from unittest import TestCase, skipUnless

from sqlalchemy import extract, func, select

from app.calculation_engine import effective_date
from app.core_routes import derived_savings, goals as goals_endpoint
from app.database import SessionLocal
from app.main import late_income_settings
from app.models import BudgetPlan, Goal, Transaction

try:
    import openpyxl
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAVE_OPENPYXL = False

XLSX_PATH = Path(os.getenv("MONEY_XLSX_PATH", "/imports/Money.xlsx"))
XLSX_AVAILABLE = HAVE_OPENPYXL and XLSX_PATH.is_file()
SKIP_REASON = f"Workbook non disponibile a {XLSX_PATH}"


def num(value) -> float:
    return round(float(value or 0), 2)


def _load_tracking_totals(year: int, month: int | None, tx_type: str) -> dict[str, float]:
    """Replica SUM(Amount WHERE Type=type AND Category=item AND YEAR/MONTH match),
    la formula reale del foglio Budget Dashboard (colonna J, riga 13+).
    Il match sulla categoria è case-insensitive, come l'operatore = di Excel.
    """
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Transactions"]
    totals: dict[str, float] = defaultdict(float)
    for row in ws.iter_rows(min_row=14, max_row=2503, min_col=3, max_col=14):
        type_, category, amount = row[1].value, row[2].value, row[3].value
        eff_date = row[11].value
        if type_ != tx_type or eff_date is None:
            continue
        if eff_date.year != year:
            continue
        if month is not None and eff_date.month != month:
            continue
        if amount is None or category is None:
            continue
        key = str(category).strip().lower()
        totals[key] += float(amount)

    canonical: dict[str, str] = {}
    bp = wb["Budget Planning"]
    for row in bp.iter_rows(min_row=1, max_row=200, min_col=1, max_col=6):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                k = cell.value.strip().lower()
                if k not in canonical:
                    canonical[k] = cell.value.strip()
    return {canonical.get(k, k): round(v, 2) for k, v in totals.items()}


def _db_budget_actual(session, year: int, month: int, budget_type: str) -> dict[str, float]:
    """Come budget_actual dell'app: aggrega ignorando maiuscole/minuscole.

    Il workbook contiene la stessa categoria con grafie diverse ("Utilities" e
    "utilities"): Excel le somma insieme perche' il confronto con = e'
    case-insensitive, quindi il confronto di parita' deve fare lo stesso.
    I template delle ricorrenze non sono movimenti e restano fuori, e con loro
    i movimenti esclusi dal budget (`counts_in_budget = false`): il workbook non
    ha quel concetto, ma la cifra che l'app mostra nel budget si', e il confronto
    deve essere fra le due cifre del budget.
    """
    rows = session.execute(
        select(Transaction.category, func.sum(Transaction.amount)).where(
            extract("year", Transaction.effective_on) == year,
            extract("month", Transaction.effective_on) == month,
            Transaction.transaction_type == budget_type,
            Transaction.is_recurring_template.is_(False),
            Transaction.counts_in_budget.is_(True),
        ).group_by(Transaction.category)
    ).all()
    totals: dict[str, float] = defaultdict(float)
    for category, total in rows:
        if total:
            totals[str(category).strip().lower()] += float(total)
    return {key: round(value, 2) for key, value in totals.items()}


@skipUnless(XLSX_AVAILABLE, SKIP_REASON)
class BudgetDashboardParityTests(TestCase):
    """Fase 3.1: Budget Dashboard (Expenses / Income / Savings)."""

    def test_expenses_august_2026(self):
        expected = _load_tracking_totals(2026, 8, "Expenses")
        with SessionLocal() as session:
            actual = _db_budget_actual(session, 2026, 8, "Expenses")
        self.assertEqual({k.strip().lower(): v for k, v in expected.items()}, {k: v for k, v in actual.items() if v})

    def test_expenses_april_2024(self):
        expected = _load_tracking_totals(2024, 4, "Expenses")
        with SessionLocal() as session:
            actual = _db_budget_actual(session, 2024, 4, "Expenses")
        self.assertEqual({k.strip().lower(): v for k, v in expected.items()}, {k: v for k, v in actual.items() if v})

    def test_income_october_2025(self):
        expected = _load_tracking_totals(2025, 10, "Income")
        with SessionLocal() as session:
            actual = _db_budget_actual(session, 2025, 10, "Income")
        self.assertEqual({k.strip().lower(): v for k, v in expected.items()}, {k: v for k, v in actual.items() if v})

    def test_savings_is_income_minus_expenses(self):
        """Il risparmio non e' piu' la somma dei movimenti di tipo Savings: e'
        quello che resta. Il confronto con l'oracolo si sposta di conseguenza,
        sulle sue stesse righe di entrate e uscite.
        """
        for year, month in ((2026, 4), (2026, 7), (2025, 10)):
            with self.subTest(period=f"{year}-{month:02d}"):
                income = sum(_load_tracking_totals(year, month, "Income").values())
                expenses = sum(_load_tracking_totals(year, month, "Expenses").values())
                with SessionLocal() as session:
                    actual = derived_savings(session, year, month)
                self.assertEqual(round(income - expenses, 2), actual)


@skipUnless(XLSX_AVAILABLE, SKIP_REASON)
class BudgetAnnualParityTests(TestCase):
    """Fase 3.1: Budget Annual (somma sui 12 mesi dell'anno)."""

    def test_expenses_2025(self):
        expected = _load_tracking_totals(2025, None, "Expenses")
        with SessionLocal() as session:
            plans = session.scalars(
                select(BudgetPlan).where(
                    extract("year", BudgetPlan.period) == 2025,
                    BudgetPlan.budget_type == "Expenses",
                )
            ).all()
            categories = {p.category for p in plans}
            actual: dict[str, float] = defaultdict(float)
            for month in range(1, 13):
                for category, total in _db_budget_actual(session, 2025, month, "Expenses").items():
                    actual[category] += total
        actual = {k: round(v, 2) for k, v in actual.items() if v}
        self.assertEqual({k.strip().lower(): v for k, v in expected.items()}, {k.strip().lower(): v for k, v in actual.items()})


@skipUnless(XLSX_AVAILABLE, SKIP_REASON)
class GoalsParityTests(TestCase):
    """Fase 3.1 (collaterale): Goals, foglio Goals righe 40+."""

    def test_active_and_completed_goal(self):
        wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
        ws = wb["Goals"]
        expected = {}
        for r in (40, 46):
            name = ws.cell(r, 4).value
            current = ws.cell(r, 17).value
            if name and current is not None:
                expected[name] = round(float(current), 2)

        # Si chiede il numero all'app, non se ne rifa' il calcolo qui: un test di
        # parita' deve confrontare l'Excel con cio' che l'utente vede. Rifacendo
        # la somma a mano, il test resta verde anche quando l'app sbaglia - e
        # diventa rosso quando l'app viene corretta, che e' esattamente quello
        # che e' successo con il verso dei movimenti taggati.
        with SessionLocal() as session:
            actual = {voce["name"]: voce["currentAmount"] for voce in goals_endpoint(session)["items"]}

        for name, exp_val in expected.items():
            with self.subTest(goal=name):
                self.assertEqual(actual.get(name), exp_val)


@skipUnless(XLSX_AVAILABLE, SKIP_REASON)
class EffectiveDateParityTests(TestCase):
    """Data di competenza: colonna N del foglio Transactions.

    Formula del workbook (identica su ogni riga della tabella Tracking):
        =IF(AND(Type="Income", shift_late_income_status="Active",
                DAY(Date)>=shift_late_income_starting_date),
            DATE(YEAR(Date), MONTH(Date)+1, 1), Date)

    Sono i due test che intercettano la divergenza piu' insidiosa: un movimento
    creato dall'app senza applicare la regola finirebbe nel mese sbagliato in
    ogni aggregato (budget, Panoramica, tasso di risparmio) senza errori visibili.
    """

    def test_every_transaction_matches_the_rule(self):
        with SessionLocal() as session:
            shift, day = late_income_settings(session)
            wrong = [
                (tx.id, tx.occurred_on, tx.effective_on, effective_date(tx.occurred_on, tx.transaction_type, shift, day))
                for tx in session.scalars(select(Transaction)).all()
                if tx.effective_on != effective_date(tx.occurred_on, tx.transaction_type, shift, day)
            ]
        self.assertEqual(wrong, [], f"{len(wrong)} movimenti con data di competenza fuori regola")

    def test_matches_workbook_column_n(self):
        wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
        ws = wb["Transactions"]
        expected: dict[int, date] = {}
        for row_number, row in enumerate(ws.iter_rows(min_row=14, max_row=2503, min_col=3, max_col=14), start=14):
            occurred, eff = row[0].value, row[11].value
            if occurred is None or eff is None:
                continue
            expected[row_number] = eff.date() if hasattr(eff, "date") else eff

        with SessionLocal() as session:
            rows = session.scalars(select(Transaction).where(Transaction.source_row.is_not(None))).all()
            mismatched = [
                (tx.source_row, tx.effective_on, expected[tx.source_row])
                for tx in rows
                if tx.source_row in expected and tx.effective_on != expected[tx.source_row]
            ]
        self.assertEqual(mismatched, [], f"{len(mismatched)} righe con effective date diversa dal workbook")
