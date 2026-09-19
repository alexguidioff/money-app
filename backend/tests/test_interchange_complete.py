from datetime import date
from decimal import Decimal as D
from io import BytesIO
from unittest import TestCase

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.database import Base
from app.interchange import build_export, SHEETS
from app.interchange_import import import_data, read_and_validate, InterchangeError
from app.models import (Account, AccountValuation, Goal, GoalMilestone, IncomeStream, LiabilityProfile,
                        RetirementProfile)


class CompleteExportTests(TestCase):
    def setUp(self):
        self.engines = [create_engine('sqlite://'), create_engine('sqlite://')]
        for engine in self.engines:
            Base.metadata.create_all(engine)
        self.source, self.target = [Session(engine) for engine in self.engines]
        self.source.add_all([
            Account(id=41, name='Casa', source_group='asset', needs_manual_valuation=True,
                    notes='=Valore da perizia', is_liquid=False),
            Account(id=42, name='Broker', source_group='bank', is_broker=True),
            Account(id=43, name='Linea', source_group='liability', starting_balance=-1000),
            AccountValuation(id=71, account_id=41, observed_on=date(2025,1,1), value=D('120000.12'), notes='Prima'),
            AccountValuation(id=72, account_id=41, observed_on=date(2026,1,1), value=D('130000.35'), notes='Seconda'),
            LiabilityProfile(id=81, account_id=43, kind='credit_line', credit_limit=D('9000.50'),
                             debt_type='revolving', original_principal=1, annual_rate=D('2.125'),
                             start_date=date(2025,1,1), end_date=date(2030,1,1), notes='Linea test'),
            RetirementProfile(id=91, birth_year=1990, country='CH', target_retirement_age=55,
                real_return=D('3.25'), withdrawal_rate=D('3.50'), withdrawal_tax_rate=D('12.50'),
                inflation=D('1.75'), expense_basis='custom', custom_annual_expenses=D('20000.15'),
                lean_annual_expenses=D('15000.01'), expense_rules='[{"category":"Mutuo","mode":"drop"}]', notes='Ipotesi'),
            IncomeStream(id=101, name='Rendita', kind='annuity', amount=D('14000.50'), start_age=65,
                         indexed=False, country='CH', amount_if_stopping_now=D('4000.25'), notes='Stima ente'),
            IncomeStream(id=102, name='Capitale', kind='capital', amount=D('70000.75'), start_age=60, country='CH'),
            # Id alti apposta: con gli id 1 e 2 la rimappatura cadrebbe sugli
            # stessi numeri e il test delle tappe passerebbe anche senza
            # rimappare nulla.
            Goal(id=201, name='Patrimonio', kind='net_worth', target_amount=500000),
            Goal(id=202, name='Riserva', kind='contributions', target_account='Broker', target_amount=10000),
        ])
        self.source.commit()

    def tearDown(self):
        self.source.close(); self.target.close()
        for engine in self.engines:
            engine.dispose()

    def test_all_data_and_remapped_account_references(self):
        import_data(self.target, build_export(self.source))
        for title in ('Conti','ProfiloPensione','FlussiPensione','Obiettivi','Debiti','ValutazioniConti'):
            model, fields = SHEETS[title]
            fields = [f for f in fields if f not in ('id','account_id')]
            before = [tuple(getattr(r,f) for f in fields) for r in self.source.scalars(select(model).order_by(model.id))]
            after = [tuple(getattr(r,f) for f in fields) for r in self.target.scalars(select(model).order_by(model.id))]
            self.assertEqual(before, after, title)
        casa = self.target.scalar(select(Account).where(Account.name=='Casa'))
        self.assertNotEqual(41, casa.id)
        self.assertTrue(all(r.account_id == casa.id for r in self.target.scalars(select(AccountValuation))))
        linea = self.target.scalar(select(Account).where(Account.name=='Linea'))
        self.assertEqual(linea.id, self.target.scalar(select(LiabilityProfile)).account_id)

    def test_version_16_still_imports_with_defined_defaults(self):
        wb = load_workbook(build_export(self.source))
        for row in wb['Meta']:
            if row[0].value == 'versione': row[1].value = '1.6'
            if row[0].value in ('righe:ValutazioniConti','righe:ProfiloPensione','righe:FlussiPensione'):
                row[0].value = None
        for sheet in ('ValutazioniConti','ProfiloPensione','FlussiPensione'):
            del wb[sheet]
        for title, fields in {'Conti':['notes','is_broker','needs_manual_valuation'],
                              'Debiti':['kind','credit_limit'], 'Obiettivi':['kind','target_account']}.items():
            for index in range(wb[title].max_column,0,-1):
                if wb[title].cell(1,index).value in fields: wb[title].delete_cols(index)
        data = BytesIO(); wb.save(data); data.seek(0)
        import_data(self.target, data)
        self.assertEqual('term_loan', self.target.scalar(select(LiabilityProfile)).kind)
        self.assertIsNone(self.target.scalar(select(RetirementProfile)))
        self.assertEqual(3,len(list(self.target.scalars(select(Account)))))

    def test_missing_sheet_or_column_in_new_export_rejected_before_replacement(self):
        for mode in ('sheet','column'):
            wb = load_workbook(build_export(self.source))
            if mode == 'sheet': del wb['FlussiPensione']
            else:
                index = [c.value for c in wb['Conti'][1]].index('notes') + 1
                wb['Conti'].delete_cols(index)
            data = BytesIO(); wb.save(data); data.seek(0)
            with self.subTest(mode=mode), self.assertRaises(InterchangeError):
                import_data(self.source,data)
            self.assertEqual(2,len(list(self.source.scalars(select(IncomeStream)))))

    def test_ogni_colonna_di_ogni_foglio_e_esportata_o_esclusa_di_proposito(self):
        # Il controllo sulle entita' nuove non vedeva i fogli vecchi: e' cosi'
        # che "incompleto accettato" sui movimenti si perdeva a ogni reimport.
        # Una colonna nuova deve finire nell'export o in questo elenco, con il
        # motivo: non puo' restare fuori per dimenticanza.
        derivate_o_storiche = {
            'Movimenti': {'source_row', 'effective_on'},   # riga Excel d'origine; ricalcolata dall'import
            'Obiettivi': {'source_row'},
            'LedgerInvestimenti': {'source_row'},
            'CollegamentiLedger': {'created_at'},          # data tecnica, rimessa alla reimportazione
            'Note': {'source_cell', 'updated_at'},
            'Impostazioni': {'id'},
        }
        for title, (model, fields) in SHEETS.items():
            colonne = set(model.__table__.columns.keys()) - {'user_id'} - derivate_o_storiche.get(title, set())
            self.assertEqual(colonne, set(fields), title)

    def test_le_tappe_seguono_l_obiettivo_rimappato(self):
        # Il goal_id e' un id interno come gli altri: reimportando in un altro
        # account, senza rimappatura la tappa finirebbe sotto l'obiettivo di
        # qualcun altro - o sotto nessuno, se quell'id li' non esiste.
        riserva = self.source.scalar(select(Goal).where(Goal.name == 'Riserva'))
        self.source.add_all([
            GoalMilestone(goal_id=riserva.id, name='Meta\' strada', target_amount=D('5000.50'),
                          target_date=date(2026, 6, 30)),
            # Una tappa senza data: e' facoltativa, e la colonna vuota deve
            # restare vuota anche dopo il giro.
            GoalMilestone(goal_id=riserva.id, name='Traguardo', target_amount=D('10000')),
        ])
        self.source.commit()

        import_data(self.target, build_export(self.source))

        obiettivo = self.target.scalar(select(Goal).where(Goal.name == 'Riserva'))
        self.assertNotEqual(riserva.id, obiettivo.id)
        tappe = self.target.scalars(select(GoalMilestone).order_by(GoalMilestone.id)).all()
        self.assertEqual(['Meta\' strada', 'Traguardo'], [t.name for t in tappe])
        self.assertTrue(all(t.goal_id == obiettivo.id for t in tappe))
        self.assertEqual([D('5000.50'), D('10000')], [t.target_amount for t in tappe])
        self.assertEqual([date(2026, 6, 30), None], [t.target_date for t in tappe])

    def test_schema_coverage_new_entities(self):
        # Una nuova ipotesi previdenziale dimenticata nell'export deve far
        # fallire il test, anche se la reimportazione dei campi vecchi passa.
        for title in ('Conti','ProfiloPensione','FlussiPensione','ValutazioniConti','Debiti'):
            model, fields = SHEETS[title]
            self.assertEqual(set(model.__table__.columns.keys()) - {'user_id'}, set(fields), title)
