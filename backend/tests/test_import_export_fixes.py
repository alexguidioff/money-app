"""Regression checks on temporary databases and in-memory uploads only."""
import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import create_engine, select, func
from sqlalchemy.orm import Session

from app.database import Base, set_current_user, reset_current_user
from app import main, backup
from app.csv_importer import CSVStatementParser
from app.pdf_importer import BankStatementParser
from app.interchange import build_export
from app.interchange_import import read_and_validate, import_data, InterchangeError
from app.models import Account, Transaction, User, InvestmentTransaction, InvestmentTransactionDetail, TransactionLedgerLink
from app.reports import excel_report, pdf_report
from tests.test_interchange_roundtrip import _populate


class ImportExportFixes(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        Base.metadata.create_all(self.engine)
        self.session = Session(self.engine)

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    def upload(self, name, content):
        return UploadFile(filename=name, file=BytesIO(content))

    def test_csv_pdf_same_amounts_dates_and_no_period_filter(self):
        csv = "Data;Descrizione;Importo\n06/09/2026;Stipendio;1.500,00\n31/08/2026;Spesa;-1.234,56\nerrata;Da verificare;50\n06/09/2026;Zero;abc"
        rows = CSVStatementParser.extract_transactions_from_csv(csv)
        pdf_rows = BankStatementParser._parse_transaction_table([r.split(';') for r in csv.splitlines()])
        self.assertEqual(rows, pdf_rows)
        self.assertEqual([r["amount"] for r in rows], [1500, -1234.56, 50])
        self.assertIsNone(rows[2]["occurredOn"])
        result = asyncio.run(main.import_csv_statement(self.upload("ESTRATTO.CSV", csv.encode()), self.session))
        self.assertEqual(result["count"], 3)
        self.assertEqual(result["transactions"][0]["transactionType"], "Income")
        with patch.object(BankStatementParser, "extract_transactions_from_pdf", return_value=pdf_rows):
            pdf = asyncio.run(main.import_pdf_statement(self.upload("ESTRATTO.PDF", b"pdf"), self.session))
        self.assertEqual(result, pdf)

    def test_upload_boundary_empty_and_unrecognized(self):
        for name, content, status in [(None, b"x", 400), ("wrong.pdf", b"x", 400), ("empty.CSV", b"", 400),
                                      ("empty.csv", b"Data;Descrizione;Importo", 422)]:
            with self.subTest(name=name), self.assertRaises(HTTPException) as caught:
                asyncio.run(main.import_csv_statement(self.upload(name, content), self.session))
            self.assertEqual(caught.exception.status_code, status)
        with patch.object(main, "MAX_UPLOAD_BYTES", 3), self.assertRaises(HTTPException) as caught:
            asyncio.run(main.read_upload(self.upload("ok.csv", b"1234"), ".csv"))
        self.assertEqual(caught.exception.status_code, 413)

    def test_partial_save_validates_money_accounts_and_duplicates(self):
        self.session.add(Account(name="Bank", source_group="bank"))
        self.session.add(Account(name="Other", source_group="bank"))
        self.session.commit()
        good = {"date": "2026-09-06", "amount": 1500, "transactionType": "Income",
                "accountName": "Bank", "description": "Salary", "category": "Salary"}
        invalids = [dict(good, date=None), dict(good, amount=-20), dict(good, amount="NaN"),
                    dict(good, accountName="Missing"), dict(good, amount=1.001),
                    dict(good, transactionType="Transfers", destinationName="Bank")]
        result = asyncio.run(main.save_pdf_transactions([good, *invalids], self.session))
        self.assertEqual(result["saved"], 1)
        self.assertEqual([e["index"] for e in result["errors"]], list(range(1, 7)))
        self.assertEqual(self.session.scalar(select(Transaction.amount)), Decimal("1500"))
        preview = main.statement_preview([{"occurredOn": good["date"], "amount": 1500, "description": "Salary"}], self.session)
        self.assertTrue(preview["transactions"][0]["duplicate"])
        transfer = dict(good, transactionType="Transfers", destinationName="Other", amount=10)
        self.assertEqual(asyncio.run(main.save_pdf_transactions([transfer], self.session))["saved"], 1)

    def test_interchange_normalizes_negative_transfer_only(self):
        self.session.add_all([Account(name="Bank", source_group="bank"),
                              Account(name="Other", source_group="bank")])
        self.session.add(Transaction(occurred_on=date(2026, 9, 1), effective_on=date(2026, 9, 1),
                                     transaction_type="Transfers", category="_", amount=Decimal("-10.00"),
                                     account_name="Bank", account_type="Bank", destination_name="Other",
                                     destination_type="Financial"))
        self.session.commit()
        _, data = read_and_validate(BytesIO(build_export(self.session).getvalue()))
        transfer = data["Movimenti"][0]
        self.assertEqual((transfer["amount"], transfer["account_name"], transfer["account_type"],
                          transfer["destination_name"], transfer["destination_type"]),
                         (Decimal("10.00"), "Other", "Financial", "Bank", "Bank"))
        self.session.scalar(select(Transaction)).transaction_type = "Expenses"
        self.session.commit()
        with self.assertRaisesRegex(InterchangeError, "Importo negativo"):
            read_and_validate(BytesIO(build_export(self.session).getvalue()))

    def test_invalid_interchange_does_not_create_backup(self):
        with patch.object(main, "create_backup") as create, self.assertRaises(HTTPException) as caught:
            asyncio.run(main.import_data_route(self.upload("broken.xlsx", b"not xlsx"), self.session))
        self.assertEqual(caught.exception.status_code, 400)
        create.assert_not_called()

    def test_missing_backup_aborts_import(self):
        _populate(self.session)
        blob = build_export(self.session).getvalue()
        before = self.session.scalar(select(func.count(Transaction.id)))
        with patch.object(main, "create_backup", side_effect=RuntimeError("disk full")), self.assertRaises(HTTPException) as caught:
            asyncio.run(main.import_data_route(self.upload("data.xlsx", blob), self.session))
        self.assertEqual(caught.exception.status_code, 503)
        self.assertEqual(self.session.scalar(select(func.count(Transaction.id))), before)

    def test_remaps_all_references_and_accepts_legacy_export(self):
        _populate(self.session)
        tx = self.session.scalar(select(Transaction).where(Transaction.is_recurring_template.is_(False)))
        ledger = self.session.scalar(select(InvestmentTransaction))
        self.session.add(TransactionLedgerLink(transaction_id=tx.id, ledger_id=ledger.id))
        self.session.commit()
        workbook = load_workbook(build_export(self.session))
        # Deliberately non-sequential IDs make accidental preservation observable.
        for title in ("Movimenti", "LedgerInvestimenti"):
            for row in workbook[title].iter_rows(min_row=2):
                row[0].value += 100
        workbook["Movimenti"]["P3"] = 101
        workbook["DettagliLedger"]["B2"] = 101
        workbook["CollegamentiLedger"]["B2"] = 102
        workbook["CollegamentiLedger"]["C2"] = 101
        blob = BytesIO()
        workbook.save(blob)
        blob.seek(0)
        import_data(self.session, blob)
        self.session.expire_all()
        child = self.session.scalar(select(Transaction).where(Transaction.is_recurring_template.is_(False)))
        parent = self.session.scalar(select(Transaction).where(Transaction.is_recurring_template.is_(True)))
        link = self.session.scalar(select(TransactionLedgerLink))
        detail = self.session.scalar(select(InvestmentTransactionDetail))
        self.assertEqual(child.recurrence_parent_id, parent.id)
        self.assertEqual(link.transaction_id, child.id)
        self.assertEqual(link.ledger_id, detail.transaction_id)
        self.assertNotEqual(child.id, 102)
        legacy = load_workbook(build_export(self.session))
        del legacy["CollegamentiLedger"]
        for row in legacy["Meta"]:
            if row[0].value == "versione": row[1].value = "1.0"
            if row[0].value == "righe:CollegamentiLedger": row[0].value = "legacy-unused"
        blob = BytesIO()
        legacy.save(blob)
        blob.seek(0)
        self.assertEqual(read_and_validate(blob)[1]["CollegamentiLedger"], [])

    def test_invalid_metadata_and_reference_fail_before_writes(self):
        _populate(self.session)
        for sheet, cell, value in [("Meta", "B6", "not a count"), ("Movimenti", "P3", 9999),
                                   ("Movimenti", "E2", "NaN"), ("Movimenti", "A3", 1)]:
            book = load_workbook(build_export(self.session))
            book[sheet][cell] = value
            blob = BytesIO()
            book.save(blob)
            blob.seek(0)
            with self.subTest(cell=cell), self.assertRaises(InterchangeError):
                read_and_validate(blob)

    def test_only_first_user_can_manage_global_backups(self):
        self.session.add_all([User(id=3, username="first", display_name="First"),
                              User(id=9, username="second", display_name="Second")])
        self.session.commit()
        token = set_current_user(9)
        try:
            with self.assertRaises(HTTPException) as caught: main.require_backup_admin(self.session)
            self.assertEqual(caught.exception.status_code, 403)
        finally: reset_current_user(token)
        token = set_current_user(3)
        try: main.require_backup_admin(self.session)
        finally: reset_current_user(token)
        for route in main.app.routes:
            if route.path.startswith("/api/backups"):
                self.assertIn(main.require_backup_admin, [d.call for d in route.dependant.dependencies])

    def test_restore_is_atomic_and_exit_one_is_failure(self):
        with TemporaryDirectory() as folder:
            target = Path(folder) / "test.dump"
            target.touch()
            with patch.object(backup, "BACKUPS_DIR", Path(folder)), patch.object(backup, "create_backup", return_value={"filename": "safety.dump"}), patch.object(
                backup.subprocess, "run", side_effect=[SimpleNamespace(returncode=0), SimpleNamespace(returncode=1, stderr=b"broken")]) as run:
                with self.assertRaises(RuntimeError): backup.restore_backup("test.dump")
                command = run.call_args.args[0]
                self.assertIn("--single-transaction", command)
                self.assertIn("--exit-on-error", command)
        # Never drop the schema before pg_restore.
        self.assertNotIn("psql", [call.args[0][0] for call in run.call_args_list])

    def test_report_currency(self):
        report = {"period": "2026-09", "currency": "CHF", "income": 1500, "expenses": 50, "savings": 1450,
                  "net_worth": 2000, "trend": [{"month": "09", "income": 1500, "expenses": 50}], "transactions": []}
        book = load_workbook(excel_report(report))
        self.assertIn("CHF", book["Riepilogo"]["B4"].number_format)
        from pypdf import PdfReader
        self.assertIn("CHF", PdfReader(pdf_report(report)).pages[0].extract_text())


if __name__ == "__main__":
    unittest.main()
