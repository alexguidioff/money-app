"""La migrazione che fa nascere le categorie dai nomi scritti nelle colonne.

I nomi e gli importi qui sono inventati e tondi: questo repository e' pubblico e
i valori veri di chi usa l'app non ci entrano.
"""
import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine, select, text
from sqlalchemy.orm import Session

from app.database import Base
from app.interchange import FORMAT_VERSION, build_export
from app.interchange_import import import_data
from app.migrations import tracked_changes
from app.models import Account, BudgetPlan, CategorizationRule, Category, Transaction


class MigrazioneCategorieTests(unittest.TestCase):
    """Un database di prima dell'albero: i nomi stanno nelle colonne.

    Le colonne stringa non esistono piu' nel modello - le toglie l'ultimo passo
    del lavoro - quindi qui si rimettono a mano, come le trova la migrazione su
    un database vero.
    """

    def setUp(self):
        self.engine = create_engine("sqlite://")
        Base.metadata.create_all(self.engine)
        with self.engine.begin() as conn:
            for tabella in ("transactions", "budget_plans", "categorization_rules"):
                conn.execute(text(f"ALTER TABLE {tabella} ADD COLUMN category VARCHAR(255)"))
            conn.execute(text("ALTER TABLE budget_plans ADD COLUMN category_group VARCHAR(255)"))
            conn.execute(text("INSERT INTO transactions (user_id, occurred_on, effective_on, transaction_type,"
                              " category, amount, details, counts_in_budget, is_recurring_template) VALUES"
                              " (1, '2026-01-10', '2026-01-10', 'Expenses', 'Vitto', 30, 'Spesa', true, false),"
                              " (1, '2026-01-11', '2026-01-11', 'Expenses', 'Vitto', 20, 'Spesa', true, false),"
                              " (1, '2026-01-12', '2026-01-12', 'Expenses', 'Trasporti', 10, 'Abbonamento', true, false),"
                              " (1, '2026-01-13', '2026-01-13', 'Transfers', '_', 40, 'Giroconto', false, false),"
                              " (1, '2026-01-14', '2026-01-14', 'Income', NULL, 50, 'Entrata', true, false)"))
            conn.execute(text("INSERT INTO budget_plans (user_id, period, budget_type, category, category_group, amount)"
                              " VALUES (1, '2026-01-01', 'Expenses', 'Casa', NULL, 100),"
                              " (1, '2026-02-01', 'Expenses', 'Casa', NULL, 100)"))
            conn.execute(text("INSERT INTO categorization_rules (user_id, position, pattern, is_regex, category, active)"
                              " VALUES (1, 1, 'spesa', false, 'Vitto', true)"))
            conn.execute(text("INSERT INTO lookup_options (user_id, option_group, position, value)"
                              " VALUES (1, 'categories_expenses', 1, 'Salute')"))
        self.session = Session(self.engine)

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    def radici(self):
        return {riga.name: riga for riga in self.session.scalars(
            select(Category).where(Category.parent_id.is_(None)))}

    def test_10_una_radice_per_nome_e_movimenti_collegati(self):
        tracked_changes(self.engine)
        self.session.expire_all()
        radici = self.radici()
        # I nomi vengono da tutti e tre i posti: movimenti, budget e vocabolario.
        self.assertEqual(set(radici), {"Vitto", "Trasporti", "Casa", "Salute"})
        movimenti = self.session.scalars(select(Transaction).order_by(Transaction.id)).all()
        self.assertEqual([riga.category_id for riga in movimenti],
                         [radici["Vitto"].id, radici["Vitto"].id, radici["Trasporti"].id, None, None])
        self.assertEqual([riga.category_id for riga in self.session.scalars(select(BudgetPlan).order_by(BudgetPlan.id))],
                         [radici["Casa"].id, radici["Casa"].id])
        self.assertEqual(self.session.scalar(select(CategorizationRule)).category_id, radici["Vitto"].id)

    def test_11_eseguita_due_volte_non_duplica_niente(self):
        tracked_changes(self.engine)
        prima = {riga.name: riga.id for riga in self.session.scalars(select(Category))}
        collegati = self.session.scalars(select(Transaction.category_id)).all()
        tracked_changes(self.engine)
        self.session.expire_all()
        self.assertEqual({riga.name: riga.id for riga in self.session.scalars(select(Category))}, prima)
        self.assertEqual(self.session.scalars(select(Transaction.category_id)).all(), collegati)

    def test_12_il_segnaposto_dei_trasferimenti_non_diventa_una_categoria(self):
        tracked_changes(self.engine)
        self.session.expire_all()
        self.assertNotIn("_", self.radici())
        trasferimento = self.session.scalar(select(Transaction).where(Transaction.transaction_type == "Transfers"))
        self.assertIsNone(trasferimento.category_id)

    def test_13_non_inventa_gerarchie_che_non_erano_dichiarate(self):
        with self.engine.begin() as conn:
            conn.execute(text("UPDATE budget_plans SET category_group = 'Necessario' WHERE category = 'Casa'"))
        tracked_changes(self.engine)
        self.session.expire_all()
        radici = self.radici()
        # Solo cio' che era scritto: il gruppo dichiarato compare come radice, e
        # 'Casa' ci sta sotto. Tutto il resto resta piatto.
        self.assertIn("Necessario", radici)
        self.assertEqual(self.session.scalar(select(Category).where(Category.name == "Casa")).parent_id,
                         radici["Necessario"].id)
        for nome in ("Vitto", "Trasporti", "Salute"):
            self.assertIsNone(radici[nome].parent_id)
        self.assertEqual(len(self.session.scalars(select(Category)).all()), 5)


class ImportFileVecchiTests(unittest.TestCase):
    """Un export fatto prima dell'albero non perde le categorie.

    Il formato promette che un file piu' vecchio resta leggibile: allora la
    colonna che portava il nome va sciolta in una categoria, non ignorata.
    """

    def setUp(self):
        self.source = create_engine("sqlite://")
        Base.metadata.create_all(self.source)
        with Session(self.source) as session:
            session.add(Account(name="Conto", source_group="bank", starting_balance=Decimal("1000")))
            vitto = Category(parent_id=None, name="Vitto")
            session.add(vitto)
            session.flush()
            session.add_all([
                Transaction(occurred_on=date(2026, 3, 1), effective_on=date(2026, 3, 1),
                            transaction_type="Expenses", category_id=vitto.id, amount=Decimal("30"),
                            account_name="Conto", details="Spesa"),
                Transaction(occurred_on=date(2026, 3, 2), effective_on=date(2026, 3, 2),
                            transaction_type="Transfers", category_id=None, amount=Decimal("40"),
                            account_name="Conto", destination_name="Conto", details="Giroconto"),
            ])
            session.commit()
            self.vecchio = self._come_un_file_vecchio(session)
        self.target = create_engine("sqlite://")
        Base.metadata.create_all(self.target)

    def tearDown(self):
        self.source.dispose()
        self.target.dispose()

    @staticmethod
    def _come_un_file_vecchio(session) -> BytesIO:
        """L'export di oggi, riscritto come lo scriveva la versione di prima."""
        workbook = load_workbook(build_export(session))
        del workbook["Categorie"]
        for riga in workbook["Meta"]:
            if riga[0].value == "versione":
                riga[1].value = "1.9"
            if riga[0].value == "righe:Categorie":
                riga[0].value = None
        for foglio, modello in (("Movimenti", Transaction), ("Budget", BudgetPlan), ("RegoleCategoria", CategorizationRule)):
            sheet = workbook[foglio]
            intestazione = [cella.value for cella in sheet[1]]
            colonna = intestazione.index("category_id") + 1
            sheet.cell(1, colonna, "category")
            for riga in range(2, sheet.max_row + 1):
                identificativo = sheet.cell(riga, intestazione.index("id") + 1).value
                categoria = session.get(modello, identificativo).category_id if identificativo else None
                nome = session.scalar(select(Category.name).where(Category.id == categoria)) if categoria else None
                sheet.cell(riga, colonna, nome)
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream

    def test_14_un_file_senza_il_foglio_delle_categorie_si_importa_comunque(self):
        with Session(self.target) as session:
            import_data(session, self.vecchio, source_name="vecchio.xlsx")
            categorie = {riga.name: riga.id for riga in session.scalars(select(Category))}
            self.assertEqual(set(categorie), {"Vitto"})
            movimenti = session.scalars(select(Transaction).order_by(Transaction.id)).all()
            self.assertEqual(movimenti[0].category_id, categorie["Vitto"])
            self.assertIsNone(movimenti[1].category_id)

    def test_15_il_formato_resta_leggibile_dalla_versione_di_prima(self):
        # Il numero maggiore non cambia: e' la promessa scritta in _read_meta.
        self.assertEqual(FORMAT_VERSION.split(".")[0], "1")
        self.assertNotEqual(FORMAT_VERSION, "1.9")
