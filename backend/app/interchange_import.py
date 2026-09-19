"""Lettura del formato di scambio: rimette nel database un file di export.

E' il gemello di :mod:`interchange`. Sostituisce la lettura del workbook
originale, che dipendeva da fogli e coordinate di celle cablate.

Due scelte di fondo:

- *sostituzione, non fusione*: il file descrive lo stato completo dell'app,
  quindi l'import svuota le entita' che il formato trasporta e le riscrive.
  Fondere richiederebbe di sapere cosa fare dei record presenti solo da una
  parte, e nessuna risposta sarebbe giusta per tutti i casi;
- *gli id vengono rimappati*: i riferimenti interni restano validi anche
  importando in un altro account, senza riutilizzare le sequenze globali.

Non vengono toccate le tabelle che il formato non trasporta: la cache dei
prezzi e dei profili (si ripopola da sola dalla fonte, ed e' indicizzata per
simbolo, quindi sopravvive al cambio di id) e le vecchie tabelle alimentate
dall'Excel.
"""

from __future__ import annotations

import warnings
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from sqlalchemy import Boolean, Date, DateTime, Integer, Numeric, delete, func, select
from sqlalchemy.orm import Session

from .categorization import categoria_da_nome
from .interchange import FORMAT_VERSION, SHEETS
from .transaction_rules import SPOSTAMENTI, TIPI_MOVIMENTO
from .models import ImportBatch

# L'ordine conta: le entita' referenziate da altre vengono scritte prima, cosi'
# il file resta leggibile anche da uno strumento che controlla i riferimenti.
# Le categorie stanno in testa: le nominano movimenti, budget e regole, e
# l'ordine inverso le cancella per ultime, quando nessuno le punta piu'.
WRITE_ORDER = ["Categorie", "Conti", "ValutazioniConti", "Debiti", "Movimenti", "RateDebiti", "Budget", "Obiettivi", "Tappe",
               "LedgerInvestimenti", "DettagliLedger", "CollegamentiLedger",
               "Strumenti", "RegoleCategoria", "Note", "Impostazioni", "Opzioni", "ProfiloPensione", "FlussiPensione",
               "Eventi", "EventiMovimenti"]

# Il nome della colonna che portava la categoria quando non era ancora una riga
# sua. Un file senza la colonna nuova ma con questa si legge lo stesso, e il
# nome si scioglie in una categoria: cosi' un export fatto prima dell'albero
# non perde le categorie. Vedi ``_read_sheet`` e ``write_imported_data``.
NOMI_CATEGORIA_LEGACY = {"Movimenti": "category", "Budget": "category", "RegoleCategoria": "category"}

TRUE_VALUES = {"true", "vero", "1", "si", "yes"}
FALSE_VALUES = {"false", "falso", "0", "no"}


class InterchangeError(ValueError):
    """Il file non e' un export valido, o non e' leggibile da questa versione."""


def _read_meta(workbook) -> dict[str, Any]:
    if "Meta" not in workbook.sheetnames:
        raise InterchangeError("Il foglio 'Meta' manca: il file non e' un export di questa app.")
    meta = {row[0]: row[1] if len(row) > 1 else None
            for row in workbook["Meta"].iter_rows(min_row=2, values_only=True) if row and row[0]}
    if str(meta.get("formato", "")).strip() != "money-interchange":
        raise InterchangeError(f"Formato non riconosciuto: {meta.get('formato')!r}.")
    version = str(meta.get("versione", "")).strip()
    # Il numero maggiore segnala un cambio incompatibile; il minore aggiunge
    # soltanto colonne, e un file piu' vecchio resta leggibile. I numeri si
    # leggono come numeri: '1.10' viene dopo '1.9', non prima.
    if not all(parte.isdigit() for parte in version.split(".")):
        raise InterchangeError(f"Versione del formato non riconosciuta: {version!r}.")
    if version.split(".")[0] != FORMAT_VERSION.split(".")[0]:
        raise InterchangeError(f"Versione del formato non compatibile: {version} (attesa {FORMAT_VERSION}).")
    return meta


def _coerce(column, value: Any) -> Any:
    """Riporta il valore letto dal foglio al tipo della colonna."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    kind = column.type
    try:
        if isinstance(kind, Boolean):
            if isinstance(value, bool):
                return value
            text_value = str(value).strip().casefold()
            if text_value in TRUE_VALUES:
                return True
            if text_value in FALSE_VALUES:
                return False
            raise ValueError(f"valore booleano non riconosciuto: {value!r}")
        if isinstance(kind, DateTime):
            if isinstance(value, datetime):
                return value
            return datetime.fromisoformat(str(value).strip())
        if isinstance(kind, Date):
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            return date.fromisoformat(str(value).strip()[:10])
        if isinstance(kind, Numeric):
            number = Decimal(str(value).strip())
            if not number.is_finite():
                raise ValueError("numero non finito")
            return number
        if isinstance(kind, Integer):
            number = Decimal(str(value).strip())
            if not number.is_finite() or number != number.to_integral_value():
                raise ValueError("numero intero richiesto")
            return int(number)
    except (ValueError, TypeError, InvalidOperation) as error:
        raise InterchangeError(f"Valore non valido per '{column.key}': {value!r} ({error})") from error
    return str(value)


def _read_sheet(workbook, title: str, model, columns: list[str], version: str = FORMAT_VERSION) -> list[dict[str, Any]]:
    if title not in workbook.sheetnames:
        raise InterchangeError(f"Il foglio '{title}' manca dal file.")
    sheet = workbook[title]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    except StopIteration:
        raise InterchangeError(f"Il foglio '{title}' e' vuoto: manca perfino l'intestazione.") from None

    additions = {"counts_in_budget": True, "refund_of_id": None, "is_active": True, "is_liquid": True,
                 "repayment_start_date": None, "planned_drawdowns": None, "grace_interest": "paid",
                 "is_classified": False, "incomplete_accepted": False}
    # I ripieghi dipendono dall'entita': "kind" non ha lo stesso significato
    # per un obiettivo e per un debito. I file 1.6 non avevano questi campi.
    legacy_additions = {
        "Conti": {"notes": None, "needs_manual_valuation": False, "is_broker": False},
        "Debiti": {"kind": "term_loan", "credit_limit": None},
        "Obiettivi": {"kind": "contributions", "target_account": None},
    }.get(title, {})
    if version in {"1.0", "1.1", "1.2", "1.3", "1.4", "1.5", "1.6"}:
        additions.update(legacy_additions)
    # Il nome della categoria al posto del suo id: i fogli scritti prima che le
    # categorie fossero una tabella. Si legge la colonna vecchia e la si tiene
    # da parte, fuori dalle colonne del modello, perche' la scrittura la
    # sciolga in una categoria vera e propria.
    nome_legacy = NOMI_CATEGORIA_LEGACY.get(title)
    posizione_legacy = (header.index(nome_legacy)
                        if nome_legacy and nome_legacy in header and "category_id" not in header else None)
    if posizione_legacy is not None:
        additions["category_id"] = None
    missing = [name for name in columns if name not in header and name not in additions]
    if missing:
        raise InterchangeError(f"Nel foglio '{title}' mancano le colonne: {', '.join(missing)}.")
    # Colonne in piu' vengono ignorate: un file scritto da una versione piu'
    # recente resta leggibile finche' il numero maggiore non cambia.
    positions = {name: header.index(name) for name in columns if name in header}
    table_columns = model.__table__.columns

    records = []
    for number, row in enumerate(rows, start=2):
        if row is None or all(cell is None for cell in row):
            continue
        try:
            records.append({name: (additions[name] if name not in positions else _coerce(table_columns[name], row[positions[name]] if positions[name] < len(row) else None))
                            for name in columns})
            if posizione_legacy is not None:
                valore = row[posizione_legacy] if posizione_legacy < len(row) else None
                records[-1][nome_legacy] = str(valore).strip() if valore is not None else None
            if title == "Conti" and "is_active" not in header:
                records[-1]["is_active"] = (records[-1].get("status") or "").strip().lower() not in {"closed", "chiusa"} and (records[-1].get("name") or "").strip().lower() != "soldi da investire"
            if title == "Conti" and "is_liquid" not in header:
                records[-1]["is_liquid"] = records[-1]["source_group"] == "bank"
        except InterchangeError as error:
            raise InterchangeError(f"{title}, riga {number}: {error}") from error
    return records


def _check_declared_counts(meta: dict[str, Any], data: dict[str, list]) -> None:
    """Il foglio Meta dichiara quante righe dovrebbero esserci: se il file e'
    stato troncato o modificato a mano, meglio accorgersene prima di svuotare
    il database."""
    for title, records in data.items():
        declared = meta.get(f"righe:{title}")
        if declared is None:
            continue
        try:
            count = int(str(declared))
        except (ValueError, TypeError) as error:
            raise InterchangeError(f"Conteggio non valido per il foglio '{title}'.") from error
        if count != len(records):
            raise InterchangeError(
                f"Il foglio '{title}' dichiara {int(declared)} righe ma ne contiene {len(records)}: "
                "il file sembra incompleto."
            )


REFERENCES = {
    "Categorie": {"parent_id": "Categorie"},
    "ValutazioniConti": {"account_id": "Conti"},
    "Debiti": {"account_id": "Conti"},
    "RateDebiti": {"liability_account_id": "Conti", "transaction_id": "Movimenti",
                    "refund_of_id": "Movimenti"},
    # La categoria di un movimento e' un id interno come gli altri: senza
    # rimappatura un file reimportato in un altro account scriverebbe id che
    # laggiu' sono di un'altra categoria, o di nessuna.
    "Movimenti": {"recurrence_parent_id": "Movimenti", "refund_of_id": "Movimenti", "category_id": "Categorie"},
    "Budget": {"category_id": "Categorie"},
    "RegoleCategoria": {"category_id": "Categorie"},
    "DettagliLedger": {"transaction_id": "LedgerInvestimenti"},
    "CollegamentiLedger": {"transaction_id": "Movimenti", "ledger_id": "LedgerInvestimenti"},
    # L'obiettivo di una tappa e' un id interno: reimportando in un altro
    # account, senza rimappatura la tappa finirebbe sotto l'obiettivo di
    # qualcun altro, o sotto nessuno.
    "Tappe": {"goal_id": "Obiettivi"},
    "EventiMovimenti": {"transaction_id": "Movimenti", "event_id": "Eventi"},
}


def read_and_validate(source: str | Path | BinaryIO) -> tuple[dict, dict]:
    """Legge e valida prima di creare backup o iniziare la sostituzione."""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(source, data_only=True, read_only=True)
    except InterchangeError:
        raise
    except Exception as error:  # noqa: BLE001
        raise InterchangeError(f"Non riesco ad aprire il file: non sembra un .xlsx ({error}).") from error
    try:
        meta = _read_meta(workbook)
        # Prima si legge tutto e si valida, poi si scrive: se il file ha un
        # problema il database resta com'era.
        # Le categorie sono diventate un foglio proprio nella 1.10: prima erano
        # il nome scritto dentro movimenti, budget e regole, e di quel nome si
        # occupa ``_read_sheet``.
        senza_albero = tuple(int(parte) for parte in str(meta["versione"]).split(".")) < (1, 10)
        data = {title: ([] if title not in workbook.sheetnames
                        and (title == "CollegamentiLedger" and str(meta["versione"]) == "1.0"
                             or title == "Categorie" and senza_albero
                             or title == "Debiti" and str(meta["versione"]) in {"1.0", "1.1", "1.2"}
                             or title == "RateDebiti" and str(meta["versione"]) in {"1.0", "1.1", "1.2", "1.3"}
                             or title in {"ValutazioniConti", "ProfiloPensione", "FlussiPensione"}
                             and str(meta["versione"]) in {"1.0", "1.1", "1.2", "1.3", "1.4", "1.5", "1.6"}
                             or title == "RegoleCategoria"
                             and str(meta["versione"]) in {"1.0", "1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7"}
                             # Gli eventi sono nati con la 1.9: un file piu'
                             # vecchio non li ha, e non per questo e' rotto.
                             or title in {"Eventi", "EventiMovimenti"}
                             and str(meta["versione"]) in {"1.0", "1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8"}
                             # Le tappe sono nate con la 1.11.
                             or title == "Tappe"
                             and str(meta["versione"]) in {"1.0", "1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8", "1.9", "1.10"})
                        else _read_sheet(workbook, title, *SHEETS[title], version=str(meta["versione"]))) for title in WRITE_ORDER}
        _check_declared_counts(meta, data)
        if len(data["ProfiloPensione"]) > 1:
            raise InterchangeError("Il file contiene piu\' di un profilo pensione.")
        if meta.get("esportato_il"):
            try:
                datetime.fromisoformat(str(meta["esportato_il"]))
            except ValueError as error:
                raise InterchangeError("Data di esportazione non valida nel foglio Meta.") from error
        ids = {}
        for title, records in data.items():
            model, columns = SHEETS[title]
            if "id" in columns:
                ids[title] = {r["id"] for r in records}
                if None in ids[title] or len(ids[title]) != len(records):
                    raise InterchangeError(f"Identificativi mancanti o duplicati in '{title}'.")
            unique_groups = [[c.name for c in constraint.columns if c.name != "user_id"]
                             for constraint in model.__table__.constraints
                             if constraint.__class__.__name__ in ("UniqueConstraint", "PrimaryKeyConstraint")]
            for keys in unique_groups:
                if keys and all(k in columns for k in keys):
                    values = [tuple(r[k] for k in keys) for r in records if all(r[k] is not None for k in keys)]
                    if len(values) != len(set(values)):
                        raise InterchangeError(f"Valori duplicati in '{title}': {', '.join(keys)}.")
            for record in records:
                for key in columns:
                    col = model.__table__.columns[key]
                    if record[key] is None and not col.nullable and col.default is None and col.server_default is None:
                        raise InterchangeError(f"Campo '{key}' obbligatorio in '{title}'.")
                if title == "Movimenti":
                    if record["transaction_type"] not in set(TIPI_MOVIMENTO):
                        raise InterchangeError("Tipo movimento non valido in 'Movimenti'.")
                    if record["amount"] < 0:
                        if record["transaction_type"] not in SPOSTAMENTI:
                            raise InterchangeError("Importo negativo non valido in 'Movimenti'.")
                        record["amount"] = -record["amount"]
                        record["account_name"], record["destination_name"] = record["destination_name"], record["account_name"]
                        record["account_type"], record["destination_type"] = record["destination_type"], record["account_type"]
                    if record["amount"] == 0:
                        raise InterchangeError("Importo zero non valido in 'Movimenti'.")
        for title, references in REFERENCES.items():
            for record in data[title]:
                for key, target in references.items():
                    if record[key] is not None and record[key] not in ids[target]:
                        raise InterchangeError(f"Riferimento '{key}' inesistente in '{title}'.")
    finally:
        workbook.close()
    return meta, data


def import_data(session: Session, source: str | Path | BinaryIO, *, source_name: str = "") -> dict[str, Any]:
    meta, data = read_and_validate(source)
    result = write_imported_data(session, meta, data, source_name=source_name)
    session.commit()
    return result


def write_imported_data(session: Session, meta: dict, data: dict, *, source_name: str = "") -> dict:
    """Scrive in una sola transazione; il chiamante decide quando fare commit."""

    for title in reversed(WRITE_ORDER):
        session.execute(delete(SHEETS[title][0]))
    session.flush()

    id_maps = {}
    for title in WRITE_ORDER:
        model = SHEETS[title][0]
        records = data[title]
        objects = []
        for original in records:
            record = {k: v for k, v in original.items() if k != "id"}
            # I file scritti prima dell'albero portano il nome della categoria
            # invece del suo id, e il nome si scioglie in una categoria. Va dopo
            # la rimappatura: qui non c'e' nessun id da rimappare.
            dal_nome = "category" in record
            nome = record.pop("category", None) if dal_nome else None
            if title == "Movimenti":
                record["effective_on"] = record["occurred_on"]
                record["recurrence_parent_id"] = None
                record["refund_of_id"] = None
                # La categoria non punta a un altro movimento: si rimappa qui,
                # come fanno gli altri fogli. Solo i riferimenti a se stesso
                # aspettano lo scarico.
                record["category_id"] = (id_maps["Categorie"][record["category_id"]]
                                         if record["category_id"] is not None else None)
            elif title == "Categorie":
                # Il padre puo' stare in una riga sotto il figlio, e il suo id
                # nuovo non esiste ancora: si riattacca dopo lo scarico.
                record["parent_id"] = None
            else:
                for key, target in REFERENCES.get(title, {}).items():
                    record[key] = id_maps[target][record[key]] if record[key] is not None else None
            if dal_nome:
                record["category_id"] = categoria_da_nome(session, nome)
            objects.append(model(**record))
        session.add_all(objects)
        session.flush()
        if "id" in SHEETS[title][1]:
            id_maps[title] = {r["id"]: obj.id for r, obj in zip(records, objects)}
        if title == "Movimenti":
            for record, obj in zip(records, objects):
                parent = record["recurrence_parent_id"]
                obj.recurrence_parent_id = id_maps[title][parent] if parent is not None else None
                refund = record["refund_of_id"]
                obj.refund_of_id = id_maps[title][refund] if refund is not None else None
        if title == "Categorie":
            for record, obj in zip(records, objects):
                parent = record["parent_id"]
                obj.parent_id = id_maps[title][parent] if parent is not None else None
    session.flush()

    # La data di competenza dipende dalle impostazioni, non dal file.
    from .main import recompute_effective_dates
    recompute_effective_dates(session)

    batch = ImportBatch(
        source_name=source_name or "money-interchange",
        source_modified_at=datetime.fromisoformat(str(meta["esportato_il"])) if meta.get("esportato_il") else None,
        transaction_count=len(data["Movimenti"]),
        account_count=len(data["Conti"]),
        budget_count=len(data["Budget"]),
    )
    session.add(batch)
    session.flush()

    return {
        "imported": True,
        "formatVersion": str(meta.get("versione")),
        "exportedAt": meta.get("esportato_il"),
        "rows": {title: len(records) for title, records in data.items()},
    }


def summarize_state(session: Session) -> dict[str, int]:
    """Conteggi correnti per entita': serve a confrontare prima e dopo."""
    return {title: session.scalar(select(func.count()).select_from(model.__table__)) or 0
            for title, (model, _) in SHEETS.items()}


__all__ = ["InterchangeError", "import_data", "summarize_state"]
