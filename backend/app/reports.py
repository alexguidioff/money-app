"""Portable exports generated from the application's database."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


GREEN = "173B33"
LIME = "D9F373"


def excel_report(report: dict[str, Any]) -> BytesIO:
    currency = str(report.get('currency', 'EUR')).upper()
    number_format = f'"{currency}" #,##0.00'
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Riepilogo"
    overview.append(["Money - Report", report["period"]])
    overview.append([])
    overview.append(["Indicatore", "Valore"])
    for label, value in (
        ("Entrate", report["income"]),
        ("Spese", report["expenses"]),
        ("Risparmi", report["savings"]),
        ("Patrimonio netto", report["net_worth"]),
    ):
        overview.append([label, value])
    overview["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    overview["B1"].font = Font(bold=True, color="FFFFFF")
    for cell in overview[1]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
    for cell in overview[3]:
        cell.fill = PatternFill("solid", fgColor=LIME)
        cell.font = Font(bold=True)
    for cell in overview["B"]:
        if cell.row > 3:
            cell.number_format = number_format
    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 22

    movements = workbook.create_sheet("Movimenti")
    movements.append(["Data", "Tipo", "Categoria", "Descrizione", "Conto", "Destinazione", "Importo"])
    for item in report["transactions"]:
        movements.append([item["date"], item["type"], item["category"], item["description"], item["account"] or "", item["destination"] or "", item["signed_amount"]])
    # Una descrizione che comincia con "=" arriva dall'estratto conto della
    # banca, cioe' la scrive un esercente: openpyxl la salverebbe come formula
    # e Excel la eseguirebbe all'apertura.
    for row in movements.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in movements[1]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(bold=True, color="FFFFFF")
    for cell in movements["G"]:
        if cell.row > 1:
            cell.number_format = number_format
    for column, width in {"A": 13, "B": 14, "C": 22, "D": 42, "E": 20, "F": 20, "G": 14}.items():
        movements.column_dimensions[column].width = width
    movements.freeze_panes = "A2"
    movements.auto_filter.ref = movements.dimensions

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def pdf_report(report: dict[str, Any]) -> BytesIO:
    currency = str(report.get('currency', 'EUR')).upper()
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph("Money - Report finanziario", styles["Title"]), Paragraph(f"Periodo: {report['period']}", styles["Normal"]), Spacer(1, 7 * mm)]
    metrics = [["Entrate", "Spese", "Risparmi", "Patrimonio netto"], [format_money(report["income"], currency), format_money(report["expenses"], currency), format_money(report["savings"], currency), format_money(report["net_worth"], currency)]]
    table = Table(metrics, colWidths=[57 * mm] * 4)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173B33")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F4F5F1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8DFDB")),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([table, Spacer(1, 8 * mm), _monthly_chart(report["trend"]), Spacer(1, 6 * mm), Paragraph("Ultimi movimenti", styles["Heading2"])])
    rows = [["Data", "Descrizione", "Categoria", "Conto", "Importo"]]
    rows.extend([[item["date"], item["description"][:48], item["category"], item["account"] or "-", format_money(item["signed_amount"], currency)] for item in report["transactions"][:12]])
    ledger = Table(rows, colWidths=[28 * mm, 76 * mm, 42 * mm, 42 * mm, 30 * mm], repeatRows=1)
    ledger.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9F373")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8DFDB")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (-1, 1), (-1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(ledger)
    doc.build(story)
    stream.seek(0)
    return stream


def _monthly_chart(trend: list[dict[str, Any]]) -> Drawing:
    drawing = Drawing(650, 190)
    drawing.add(String(0, 174, "Entrate e spese - ultimi 7 mesi", fontName="Helvetica-Bold", fontSize=12, fillColor=colors.HexColor("#173B33")))
    chart = VerticalBarChart()
    chart.x = 35
    chart.y = 22
    chart.width = 575
    chart.height = 130
    chart.data = [[float(item["income"]) for item in trend], [float(item["expenses"]) for item in trend]]
    chart.categoryAxis.categoryNames = [item["month"] for item in trend]
    chart.categoryAxis.labels.fontSize = 8
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueStep = max(100, int(max([1, *[max(float(item["income"]), float(item["expenses"])) for item in trend]]) / 4))
    chart.bars[0].fillColor = colors.HexColor("#47A889")
    chart.bars[1].fillColor = colors.HexColor("#EF8E72")
    chart.groupSpacing = 10
    chart.barSpacing = 2
    drawing.add(chart)
    drawing.add(String(616, 145, "Entrate", fontSize=8, fillColor=colors.HexColor("#47A889")))
    drawing.add(String(616, 132, "Spese", fontSize=8, fillColor=colors.HexColor("#EF8E72")))
    return drawing


def format_money(value: float, currency: str = 'EUR') -> str:
    return f"{currency} {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
