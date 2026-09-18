"""Formato di scambio: un foglio per entita', nessuna formula.

E' il formato che l'app usa per esportare e reimportare i propri dati, e
sostituisce la lettura del workbook originale con i suoi range di celle
cablati. Regole del contratto:

- un foglio per entita', con intestazioni testuali stabili sulla riga 1;
- date e importi scritti come testo ISO (``2026-08-31``) e numeri semplici,
  per non dipendere dalle impostazioni locali di chi apre il file;
- nessuna formula, nessuna formattazione: il file trasporta dati, non calcoli;
- restano fuori solo i campi che l'app sa ricostruire da se' (la data di
  competenza, che dipende dalle impostazioni dello shift entrate) e le
  coordinate del vecchio workbook, che con esso spariscono. Tutto il resto
  viene trasportato, anche quando somiglia a un dato calcolato: il saldo
  progressivo dei movimenti, per esempio, nessuno lo ricalcola;
- il foglio ``Meta`` dichiara versione del formato e istante di esportazione.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (Account, AccountValuation, AppSetting, BudgetPlan, CategorizationRule, Goal, IncomeStream,
                     RetirementProfile, InvestmentInstrument, InvestmentTransaction,
                     InvestmentTransactionDetail,
                     LiabilityProfile, LiabilityTransactionDetail, LookupOption, Note, Transaction, TransactionLedgerLink)

FORMAT_VERSION = "1.8"


def _cell(value: Any) -> Any:
    """Normalizza un valore per la scrittura: date ISO, decimali come float."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "true" if value else "false"
    return value


SHEETS: dict[str, tuple[Any, list[str]]] = {
    "Conti": (Account, ["id", "source_group", "name", "starting_balance", "current_balance", "status",
                        "counts_in_net_worth", "is_active", "is_liquid", "notes", "needs_manual_valuation", "is_broker"]),
    "ValutazioniConti": (AccountValuation, ["id", "account_id", "observed_on", "value", "notes"]),
    "Debiti": (LiabilityProfile, ["id", "account_id", "debt_type", "original_principal", "annual_rate",
                                   "rate_type", "payment_frequency", "payment_structure", "start_date",
                                   "repayment_start_date", "end_date", "planned_drawdowns", "grace_interest",
                                   "status", "notes", "kind", "credit_limit"]),
    "RateDebiti": (LiabilityTransactionDetail, ["id", "liability_account_id", "transaction_id",
                                                 "refund_of_id", "kind", "principal_amount",
                                                 "interest_amount", "is_classified"]),
    "Movimenti": (Transaction, ["id", "occurred_on", "transaction_type", "category", "amount", "account_type",
                                "account_name", "destination_type", "destination_name", "goal", "details",
                                "balance", "is_recurring_template", "recurrence_rule", "recurrence_end_date",
                                "recurrence_parent_id", "counts_in_budget", "refund_of_id", "incomplete_accepted"]),
    "Budget": (BudgetPlan, ["id", "period", "budget_type", "category_group", "category", "amount"]),
    "Obiettivi": (Goal, ["id", "name", "starting_amount", "target_amount", "start_date", "target_date", "completed_at", "kind", "target_account"]),
    "ProfiloPensione": (RetirementProfile, ["id", "birth_year", "country", "target_retirement_age",
        "real_return", "return_volatility", "withdrawal_rate", "withdrawal_tax_rate", "inflation", "expense_basis",
        "custom_annual_expenses", "lean_annual_expenses", "expense_rules", "notes"]),
    "FlussiPensione": (IncomeStream, ["id", "name", "kind", "amount", "start_age", "indexed",
        "country", "amount_if_stopping_now", "notes"]),
    "LedgerInvestimenti": (InvestmentTransaction, ["id", "occurred_on", "ticker", "name", "transaction_type",
                                                   "amount", "units", "price", "currency"]),
    "DettagliLedger": (InvestmentTransactionDetail, ["id", "transaction_id", "fee", "notes"]),
    "CollegamentiLedger": (TransactionLedgerLink, ["id", "transaction_id", "ledger_id"]),
    "Strumenti": (InvestmentInstrument, ["id", "name", "provider_symbol", "isin", "asset_class", "area", "sector",
                                         "currency", "target_weight"]),
    "RegoleCategoria": (CategorizationRule, ["id", "position", "pattern", "is_regex", "category",
                                              "transaction_type", "min_amount", "max_amount", "active"]),
    "Note": (Note, ["id", "section", "title", "body", "status"]),
    "Impostazioni": (AppSetting, ["key", "label", "value"]),
    "Opzioni": (LookupOption, ["id", "option_group", "position", "value"]),
}


def build_export(session: Session) -> BytesIO:
    """Scrive l'intero contenuto dell'app nel formato di scambio."""
    workbook = Workbook()
    meta = workbook.active
    meta.title = "Meta"
    meta.append(["chiave", "valore"])
    meta.append(["formato", "money-interchange"])
    meta.append(["versione", FORMAT_VERSION])
    meta.append(["esportato_il", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    meta.append(["note", "Date in formato ISO (AAAA-MM-GG). Non inserire formule: il file trasporta dati."])

    for title, (model, columns) in SHEETS.items():
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        order = getattr(model, "id", None)
        query = select(model).order_by(order) if order is not None else select(model)
        for row_index, row in enumerate(session.scalars(query).all(), start=2):
            # Un appunto che inizia con '=' resta testo: una formula Excel
            # perderebbe il valore al successivo import con data_only=True.
            for column_index, column in enumerate(columns, start=1):
                cell = sheet.cell(row_index, column_index, _cell(getattr(row, column, None)))
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        meta.append([f"righe:{title}", sheet.max_row - 1])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
